"""Generate editable plan from the maintained Markdown source using python-docx.

Run with a Python environment containing python-docx. PDF export is a separate
native Word/LibreOffice step so the PDF has the same layout as the editable file.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
source=(ROOT/'docs/implementation_plan.md').read_text(encoding='utf-8')
document=Document()
section=document.sections[0]
section.page_width=Inches(8.27); section.page_height=Inches(11.69)
section.top_margin=section.bottom_margin=Inches(.8)
section.left_margin=section.right_margin=Inches(.85)
styles=document.styles
for name in ['Normal','Body Text','List Paragraph']:
    styles[name].font.name='Calibri'
    styles[name].font.size=Pt(11)
    styles[name].paragraph_format.space_after=Pt(7)
    styles[name].paragraph_format.line_spacing=1.12
for name,size in [('Title',30),('Heading 1',18),('Heading 2',13)]:
    styles[name].font.name='Calibri'
    styles[name].font.size=Pt(size)
    styles[name].font.color.rgb=RGBColor.from_string('000000' if name=='Title' else '123C49')
    styles[name].paragraph_format.keep_with_next=True
    styles[name].paragraph_format.space_before=Pt(16)
    styles[name].paragraph_format.space_after=Pt(8)
document.core_properties.title='Klicky implementation and research plan'
document.core_properties.subject='Accessible browser assistant prototype and staged HCI research roadmap'
document.core_properties.author='Klicky HCI project team'
document.add_paragraph('Klicky implementation\nand research plan',style='Title')
document.add_paragraph('University HCI project',style='Subtitle')
document.add_paragraph('25 October 2026')
document.add_paragraph('This plan explains the implemented browser-search prototype, the remaining live-model acceptance gate, and the path toward personalized, documentation-grounded interface guidance.')
document.add_paragraph('Prepared for the project team and university reviewers. It separates implemented capabilities, observed test evidence and planned research.')
decision=document.add_paragraph('Current decision')
decision.runs[0].bold=True
document.add_paragraph('Validate the local vision-model demonstration before expanding automation. Preserve user control and access-need-conditioned guidance as the central design direction.')
document.add_page_break()
document.add_paragraph('Contents',style='Title')
p=document.add_paragraph()
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u')
p._p.append(field)
document.add_page_break()
lines=source.splitlines()
first=next(i for i,line in enumerate(lines) if line.startswith('## '))
i=first
while i<len(lines):
    line=lines[i].strip()
    if not line: i+=1; continue
    if line.startswith('## '):
        p=document.add_paragraph(line[3:],style='Heading 1')
    elif line.startswith('### '):
        p=document.add_paragraph(line[4:],style='Heading 2')
    elif line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            row=[c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+',c) for c in row): rows.append(row)
            i+=1
        table=document.add_table(rows=1,cols=len(rows[0]))
        table.style='Light Shading Accent 1'
        for j,text in enumerate(rows[0]): table.rows[0].cells[j].text=text
        header=OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(header)
        for row in rows[1:]:
            for cell,text in zip(table.add_row().cells,row): cell.text=text
        for row in table.rows:
            props=row._tr.get_or_add_trPr(); props.append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after=Pt(5)
                    for run in paragraph.runs: run.font.size=Pt(10)
        document.add_paragraph()
        continue
    else:
        p=document.add_paragraph(line)
        p.paragraph_format.widow_control=True
    i+=1
footer=section.footer.paragraphs[0]
footer.alignment=2
footer.add_run('Klicky  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for run in footer.runs: run.font.size=Pt(9)
settings=document.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
# Remove inherited template rules, including Word's default blue Title border.
for node in document.styles.element.xpath('.//w:pBdr'):
    node.getparent().remove(node)
for node in document.element.xpath('.//w:pBdr'):
    node.getparent().remove(node)
(ROOT/'deliverables').mkdir(exist_ok=True)
document.save(ROOT/'deliverables/CUA_Implementation_Plan.docx')
print('Created editable implementation plan')
